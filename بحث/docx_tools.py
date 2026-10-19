# -*- coding: utf-8 -*-
"""أدوات بناء مستند وورد عربي (RTL) أنيق لبحث العادات الجاهلية."""
import copy
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ---------- الألوان ----------
GREEN = RGBColor(0x1E, 0x5B, 0x4F)      # أخضر داكن
GOLD = RGBColor(0x9A, 0x74, 0x1E)       # ذهبي داكن
DARK = RGBColor(0x22, 0x22, 0x22)
GRAY = RGBColor(0x5A, 0x5A, 0x5A)
MAROON = RGBColor(0x7B, 0x2D, 0x26)

HEX_GREEN = "1E5B4F"
HEX_GOLD = "9A741E"
HEX_CREAM = "FAF6EA"
HEX_PANEL = "F2F6F3"
HEX_SHAMELA = "EEF3F7"
HEX_LIGHT_GOLD = "F6EEDD"
HEX_TABLE_HEAD = "1E5B4F"
HEX_TABLE_ALT = "F4F7F5"

AR_FONT = "Sakkal Majalla"
AR_FONT_QUOTE = "Traditional Arabic"

# ---------- أدوات XML ----------
def _el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn("w:" + k), v)
    return e

def set_bidi(par, val="1"):
    pPr = par._p.get_or_add_pPr()
    for old in pPr.findall(qn("w:bidi")):
        pPr.remove(old)
    pPr.append(_el("w:bidi", val=val))

def shade_par(par, fill):
    pPr = par._p.get_or_add_pPr()
    for old in pPr.findall(qn("w:shd")):
        pPr.remove(old)
    pPr.append(_el("w:shd", val="clear", color="auto", fill=fill))

def par_borders(par, sides=("right",), color=HEX_GOLD, sz="18", space="6", val="single"):
    pPr = par._p.get_or_add_pPr()
    pBdr = pPr.find(qn("w:pBdr"))
    if pBdr is None:
        pBdr = OxmlElement("w:pBdr")
        pPr.append(pBdr)
    for side in sides:
        tag = "w:" + side
        for old in pBdr.findall(qn(tag)):
            pBdr.remove(old)
        b = OxmlElement(tag)
        b.set(qn("w:val"), val)
        b.set(qn("w:sz"), sz)
        b.set(qn("w:space"), space)
        b.set(qn("w:color"), color)
        pBdr.append(b)

def run_font(run, name=AR_FONT, size=None, bold=None, italic=None, color=None, rtl=True):
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        rFonts.set(qn(a), name)
    if rtl:
        if rPr.find(qn("w:rtl")) is None:
            rPr.append(_el("w:rtl", val="1"))
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color

def spacing(par, before=0, after=0, line=None, rule=None):
    pf = par.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if line is not None:
        pf.line_spacing = line

def align_rtl(par, align="justify"):
    m = {
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "left": WD_ALIGN_PARAGRAPH.LEFT,
    }
    par.alignment = m[align]

def keep_with_next(par, val=True):
    pf = par.paragraph_format
    pf.keep_with_next = val

# ---------- المستند ----------
def new_document():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.3)
    sec.bottom_margin = Cm(2.0)
    sec.right_margin = Cm(2.6)
    sec.left_margin = Cm(2.0)
    sec.header_distance = Cm(1.0)
    sec.footer_distance = Cm(1.0)

    st = doc.styles["Normal"]
    st.font.name = AR_FONT
    st.font.size = Pt(13)
    rpr = st.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        rf.set(qn(a), AR_FONT)
    pf = st.paragraph_format
    pf.space_after = Pt(4)
    pf.line_spacing = 1.35

    # أنماط العناوين (مبنية على Heading لضمان عمل الفهرس التلقائي)
    def mk(name, base, size, color, bold=True):
        s = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        s.base_style = doc.styles[base]
        s.font.name = AR_FONT
        s.font.size = Pt(size)
        s.font.bold = bold
        s.font.color.rgb = color
        rp = s.element.get_or_add_rPr()
        rf = OxmlElement("w:rFonts")
        for a in ("w:ascii", "w:hAnsi", "w:cs"):
            rf.set(qn(a), AR_FONT)
        rp.append(rf)
        return s

    mk("H1AR", "Heading 1", 20, GREEN)
    mk("H2AR", "Heading 2", 16.5, GREEN)
    mk("H3AR", "Heading 3", 14, GOLD)
    mk("H4AR", "Heading 4", 13, MAROON)
    return doc

def add_par(doc, style=None):
    p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    set_bidi(p)
    return p

def txt(par, text, **kw):
    r = par.add_run(text)
    run_font(r, **kw)
    return r

def page_break(doc):
    p = add_par(doc)
    r = p.add_run()
    r.add_break(WD_BREAK.PAGE)
    return p

def add_field(par, instr):
    r1 = par.add_run(); f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin"); r1._r.append(f1)
    r2 = par.add_run(); it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr; r2._r.append(it)
    r3 = par.add_run(); f3 = OxmlElement("w:fldChar"); f3.set(qn("w:fldCharType"), "separate"); r3._r.append(f3)
    r4 = par.add_run("…")
    r5 = par.add_run(); f5 = OxmlElement("w:fldChar"); f5.set(qn("w:fldCharType"), "end"); r5._r.append(f5)
    for rr in (r1, r2, r3, r4, r5):
        run_font(rr, size=11, color=GRAY)
    return par

def header_footer(doc, header_text, footer_text):
    sec = doc.sections[0]
    sec.different_first_page_header_footer = True

    hp = sec.header.paragraphs[0]
    set_bidi(hp); hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    txt(hp, header_text, size=10, color=GRAY)
    par_borders(hp, sides=("bottom",), color=HEX_GOLD, sz="6", space="2")

    fp = sec.footer.paragraphs[0]
    set_bidi(fp); fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par_borders(fp, sides=("top",), color=HEX_GOLD, sz="6", space="2")
    txt(fp, footer_text + "  —  ص ", size=10, color=GRAY)
    add_field(fp, "PAGE")
    txt(fp, " من ", size=10, color=GRAY)
    add_field(fp, "NUMPAGES")

# ---------- عناصر المحتوى ----------
def h1_banner(doc, kicker, title):
    p = add_par(doc, "H1AR")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    shade_par(p, HEX_GREEN)
    spacing(p, before=6, after=6, line=1.2)
    txt(p, kicker + "\n", size=13, bold=False, color=RGBColor(0xE8, 0xD9, 0xA8))
    txt(p, title, size=21, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
    keep_with_next(p)
    return p

def h2(doc, text):
    p = add_par(doc, "H2AR")
    spacing(p, before=14, after=6, line=1.2)
    txt(p, text, size=16.5, bold=True, color=GREEN)
    par_borders(p, sides=("bottom",), color=HEX_GOLD, sz="10", space="3")
    keep_with_next(p)
    return p

def h3(doc, text):
    p = add_par(doc, "H3AR")
    spacing(p, before=10, after=3, line=1.2)
    txt(p, text, size=14, bold=True, color=GOLD)
    keep_with_next(p)
    return p

def h4(doc, text):
    p = add_par(doc, "H4AR")
    spacing(p, before=8, after=3, line=1.2)
    txt(p, text, size=13, bold=True, color=MAROON)
    keep_with_next(p)
    return p

def body(doc, text, first_bold=None):
    p = add_par(doc)
    align_rtl(p, "justify")
    spacing(p, after=5, line=1.4)
    if first_bold:
        txt(p, first_bold, bold=True, color=GREEN)
    txt(p, text, size=13)
    return p

def bullet(doc, text, lead=None):
    p = add_par(doc)
    align_rtl(p, "justify")
    p.paragraph_format.left_indent = Cm(0.9)
    p.paragraph_format.first_line_indent = Cm(-0.5)
    spacing(p, after=3, line=1.35)
    txt(p, "•  ", size=13, color=GOLD, bold=True)
    if lead:
        txt(p, lead, bold=True, size=13)
    txt(p, text, size=13)
    return p

def quote(doc, text, source=None):
    p = add_par(doc)
    align_rtl(p, "justify")
    shade_par(p, HEX_CREAM)
    par_borders(p, sides=("right",), color=HEX_GOLD, sz="20", space="6")
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.right_indent = Cm(0.3)
    spacing(p, before=5, after=5, line=1.5)
    txt(p, "﴿ ", size=15, color=GOLD, bold=True)
    txt(p, text, size=14.5, name=AR_FONT_QUOTE, color=RGBColor(0x1A, 0x1A, 0x1A))
    txt(p, " ﴾", size=15, color=GOLD, bold=True)
    if source:
        q = add_par(doc)
        align_rtl(q, "left")
        q.paragraph_format.left_indent = Cm(0.8)
        shade_par(q, HEX_CREAM)
        spacing(q, before=0, after=7, line=1.15)
        txt(q, "— " + source, size=11.5, italic=True, color=GRAY)
    return p

def verse(doc, text, source=None):
    p = add_par(doc)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    shade_par(p, HEX_LIGHT_GOLD)
    par_borders(p, sides=("right", "left"), color=HEX_GOLD, sz="8", space="4")
    spacing(p, before=5, after=2, line=1.6)
    txt(p, text, size=14.5, name=AR_FONT_QUOTE, bold=False, color=RGBColor(0x2A, 0x20, 0x08))
    if source:
        q = add_par(doc)
        q.alignment = WD_ALIGN_PARAGRAPH.CENTER
        shade_par(q, HEX_LIGHT_GOLD)
        par_borders(q, sides=("right", "left"), color=HEX_GOLD, sz="8", space="4")
        spacing(q, before=0, after=7, line=1.15)
        txt(q, "— " + source, size=11, italic=True, color=GRAY)
    return p

def labeled(doc, label, text, fill=None, label_color=GREEN, size=12.5):
    p = add_par(doc)
    align_rtl(p, "justify")
    if fill:
        shade_par(p, fill)
        par_borders(p, sides=("right",), color=HEX_GOLD, sz="8", space="4")
        p.paragraph_format.right_indent = Cm(0.2)
    spacing(p, before=3, after=4, line=1.35)
    txt(p, label + " ", bold=True, color=label_color, size=size)
    txt(p, text, size=size)
    return p

def note(doc, text):
    p = add_par(doc)
    align_rtl(p, "justify")
    shade_par(p, HEX_PANEL)
    spacing(p, before=3, after=6, line=1.3)
    p.paragraph_format.left_indent = Cm(0.4)
    p.paragraph_format.right_indent = Cm(0.4)
    txt(p, "تنبيه: ", bold=True, size=11.5, color=MAROON)
    txt(p, text, size=11.5, color=RGBColor(0x33, 0x33, 0x33))
    return p

def set_cell(cell, text, bold=False, size=11.5, color=None, fill=None, align="right", name=AR_FONT):
    cell.text = ""
    p = cell.paragraphs[0]
    set_bidi(p)
    p.alignment = {"right": WD_ALIGN_PARAGRAPH.RIGHT, "center": WD_ALIGN_PARAGRAPH.CENTER,
                   "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}[align]
    spacing(p, before=1, after=1, line=1.2)
    txt(p, text, size=size, bold=bold, color=color, name=name)
    if fill:
        tcPr = cell._tc.get_or_add_tcPr()
        tcPr.append(_el("w:shd", val="clear", color="auto", fill=fill))
    va = OxmlElement("w:vAlign"); va.set(qn("w:val"), "center")
    cell._tc.get_or_add_tcPr().append(va)

def table(doc, headers, rows, widths=None, header_fill=HEX_TABLE_HEAD, font=11, alt=True):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = t._tbl.tblPr
    bidi = OxmlElement("w:bidiVisual")
    tblPr.append(bidi)
    hdr = t.rows[0].cells
    for i, h in enumerate(headers):
        set_cell(hdr[i], h, bold=True, size=font + 0.5, color=RGBColor(0xFF, 0xFF, 0xFF),
                 fill=header_fill, align="center")
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        fill = HEX_TABLE_ALT if (alt and ri % 2 == 1) else None
        for ci, val in enumerate(row):
            set_cell(cells[ci], str(val), size=font, fill=fill,
                     align="center" if ci == 0 else "right")
    if widths:
        for ci, w in enumerate(widths):
            for row in t.rows:
                row.cells[ci].width = Cm(w)
    spacing(doc.add_paragraph(), after=4)
    return t

def toc_field(doc, title="فهرس الموضوعات"):
    h = add_par(doc)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    spacing(h, before=4, after=10)
    txt(h, title, size=19, bold=True, color=GREEN)
    par_borders(h, sides=("bottom",), color=HEX_GOLD, sz="10", space="4")
    p = add_par(doc)
    set_bidi(p)
    add_field(p, 'TOC \\o "1-2" \\h \\z \\u')
    q = add_par(doc)
    set_bidi(q)
    spacing(q, before=4, after=2)
    txt(q, "(بعد فتح الملف في Word: اضغط داخل المستند Ctrl+A ثم F9 لتحديث الفهرس وأرقام الصفحات)",
        size=10.5, italic=True, color=GRAY)

# ---------- قالب الشاهد الموحد ----------
def witness(doc, idx, tag, title, text, cite, story=None, face=None,
            shamela=None, caveat=None, verses=False):
    """شاهد كامل: عنوان + نص حرفي (اقتباس أو شعر) + توثيق + قصة + وجه الاستشهاد + عبارة بحث + تنبيه."""
    h3(doc, f"{idx} {tag} — {title}")
    lines = text if isinstance(text, (list, tuple)) else [text]
    if verses:
        for i, ln in enumerate(lines):
            verse(doc, ln, source=(cite if i == len(lines) - 1 else None))
    else:
        quote(doc, "\n".join(lines), source=cite)
    labeled(doc, "التوثيق:", cite, fill=HEX_PANEL)
    if story:
        labeled(doc, "سياق الخبر والقصة:", story, fill=HEX_LIGHT_GOLD, label_color=MAROON)
    if face:
        labeled(doc, "وجه الاستشهاد:", face)
    if shamela:
        labeled(doc, "عبارة البحث في الشاملة:", "«" + shamela + "»", fill=HEX_SHAMELA, label_color=GREEN)
    if caveat:
        note(doc, caveat)
    return

def balance(doc, text):
    """فقرة موازنة الشواهد ودرجة الدلالة في ختام كل عادة."""
    labeled(doc, "موازنة الشواهد ودرجة الدلالة:", text, fill=HEX_PANEL, label_color=MAROON, size=12.5)
